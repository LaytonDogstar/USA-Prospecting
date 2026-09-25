"""Write the client-profile review workbook."""
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .importers import unknown_aliases
from .profiles import TIERED_CATEGORIES, build_profiles, segment_summary
from .research import fact_rows, research_by_buyer

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
EDIT_FILL = PatternFill("solid", fgColor="FFF9E6")

BUYER_COLUMNS = [
    # (header, key, number format, width, editable)
    ("Buyer", "buyer", None, 30, False),
    ("Category", "category", None, 9, True),
    ("Status", "status", None, 18, True),
    ("Website", "website", None, 26, True),
    ("Accepted leads", "accepted", "#,##0", 10, False),
    ("Revenue ($)", "revenue", "#,##0", 11, False),
    ("EPL ($)", "epl", "#,##0.00", 9, False),
    ("Volume tier", "volume_tier", None, 8, False),
    ("Price tier", "price_tier", None, 8, False),
    ("Reference segment", "reference_segment", None, 11, False),
    ("Price-reject share", "price_reject_share", "0%", 9, False),
    ("Redirect rate", "redirect_rate", "0%", 9, False),
    ("Tiers", "tiers", None, 6, False),
    ("Active tiers", "active_tiers", None, 7, False),
    ("States accepted (#)", "n_states", None, 9, False),
    ("States accepted", "states", None, 40, False),
    ("Loan min ($)", "loan_min", "#,##0", 9, False),
    ("Loan max ($)", "loan_max_label", "#,##0", 9, False),
    ("Min monthly income ($)", "income_min", "#,##0", 10, False),
    ("Min age", "age_min", "0", 7, False),
    ("Income types", "income_types", None, 26, False),
    ("Pay frequency", "pay_freq", None, 26, False),
    ("Direct deposit only", "dd_label", None, 9, False),
    ("Filters on file", "filters_label", None, 8, False),
    ("Company type (research)", "r_company_type", None, 16, False),
    ("Tribe (research)", "r_tribe", None, 22, False),
    ("Products (research)", "r_products", None, 26, False),
    ("Parent / servicer (research)", "r_parent_or_servicer", None, 26, False),
    ("Related brands (research)", "r_related_brands", None, 30, False),
    ("HQ (research)", "r_hq_location", None, 18, False),
    ("Size signals (research)", "r_size_signals", None, 30, False),
    ("Litigation / regulatory (research)", "r_litigation_or_regulatory", None, 40, False),
    ("Research confidence", "research_confidence", None, 10, False),
    ("Research notes", "research_notes", None, 40, False),
    ("Notes", "notes", None, 40, True),
]


def _flatten(p, research):
    states = p.get("states_accepted")
    r = research.get(p["buyer"], {})
    return {
        **r,
        **p,
        "website": p.get("website") or r.get("research_website"),
        "n_states": len(states) if states else None,
        "states": ", ".join(states) if states else None,
        "loan_max_label": "no limit" if p.get("loan_max_unbounded") else p.get("loan_max"),
        "dd_label": {True: "yes", False: "no", None: None}[p.get("direct_deposit_only")],
        "filters_label": "yes" if p.get("has_filters") else "no",
    }


def _sheet(wb, title, columns, rows, freeze="B2"):
    ws = wb.create_sheet(title)
    for ci, (header, _key, _fmt, width, editable) in enumerate(columns, 1):
        c = ws.cell(row=1, column=ci, value=header)
        c.font = Font(bold=True)
        c.fill = EDIT_FILL if editable else HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(ci)].width = width
    for ri, row in enumerate(rows, 2):
        for ci, (_h, key, fmt, _w, editable) in enumerate(columns, 1):
            c = ws.cell(row=ri, column=ci, value=row.get(key))
            if fmt:
                c.number_format = fmt
            if editable:
                c.fill = EDIT_FILL
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


def _readme(wb, period, counts):
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Client profile review", True),
        (f"Generated {date.today().isoformat()} from the revenue export ({period}) and the tier filter export.", False),
        ("", False),
        ("What to check", True),
        ("1. Buyers: the Category, Status, Website and Notes columns (shaded) are yours to correct. "
         "Changes go into data/reference/buyers.csv, then re-run the build.", False),
        ("2. Reference profile: what your best buyers look like. This becomes the basis for the lookalike rules in step 2.", False),
        ("3. Win-back: direct buyers with no sales this period. Already integrated, so cheaper to reactivate than a new client.", False),
        ("4. Data issues: anything the import couldn't resolve.", False),
        ("", False),
        ("How buyers are valued (option 3: volume and price)", True),
        ("Only direct lenders (D) and service providers (S) are tiered. Networks (N) and non-lender offers (O) are excluded.", False),
        ("For each measure, buyers are sorted largest first and a running total kept. Tier A = buyers while the running "
         "total is within 80% of the total, B = within 95%, C = the rest.", False),
        ("Volume tier uses accepted leads (Total Sold). Price tier uses revenue (Commission).", False),
        ("Reference segment: 'volume' = A on volume, 'price' = A on revenue, 'both' = A on both. "
         "These are the buyers new targets will be compared against.", False),
        ("", False),
        ("Assumptions (from the questions doc)", True),
        ("Revenue is gross for the last 30 days. Price Reject Tier revenue counts as normal revenue for the buyer. Tree column ignored.", False),
        ("Accepted states = all 50 states + DC minus the tier's State Blacklist; a tier with no blacklist accepts every state. "
         "A buyer's states, loan range, income and age are the loosest values across its tiers.", False),
        ("'X - 0.00' ranges mean X or more with no upper limit. Tiers whose name was a filter value in the export are treated as unnamed.", False),
        ("Filters exist for some buyers only; profile columns are blank where no filters were exported.", False),
        ("", False),
        (f"Loaded: {counts['revenue_rows']} revenue tier rows ({counts['duplicates']} exact duplicates skipped), "
         f"{counts['filter_rows']} filter rows, {counts['buyers']} buyers.", False),
        ("This file contains confidential commercial data. Do not share outside Dogstar.", True),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)


def write_review(conn: sqlite3.Connection, path: Path, period: str, duplicates: int = 0) -> dict:
    research = research_by_buyer(conn)
    profiles = [_flatten(p, research) for p in build_profiles(conn)]
    tiered = [p for p in profiles if p["category"] in TIERED_CATEGORIES]
    counts = {
        "revenue_rows": conn.execute("SELECT COUNT(*) FROM tier_revenue").fetchone()[0],
        "filter_rows": conn.execute("SELECT COUNT(*) FROM tier_filter").fetchone()[0],
        "buyers": len(profiles),
        "duplicates": duplicates,
    }

    wb = Workbook()
    _readme(wb, period, counts)
    _sheet(wb, "Buyers", BUYER_COLUMNS, profiles)

    seg_cols = [
        ("Segment", "segment", None, 36, False),
        ("Buyers", "buyers", None, 7, False),
        ("Accepted leads", "accepted", "#,##0", 10, False),
        ("Revenue ($)", "revenue", "#,##0", 11, False),
        ("Median EPL ($)", "median_epl", "#,##0.00", 9, False),
        ("Median loan min ($)", "median_loan_min", "#,##0", 9, False),
        ("Median loan max ($)", "median_loan_max", "#,##0", 9, False),
        ("Median min income ($)", "median_income_min", "#,##0", 10, False),
        ("Median states accepted", "median_states", "0", 9, False),
        ("States accepted by half or more", "states_accepted_by_half_or_more", None, 40, False),
        ("Buyers with filters", "with_filters", None, 8, False),
        ("Names", "names", None, 60, False),
    ]
    _sheet(wb, "Reference profile", seg_cols, segment_summary(profiles))

    winback = [p for p in tiered if not p["accepted"]]
    winback.sort(key=lambda p: -(p["max_sent"] or 0))
    wb_cols = [c for c in BUYER_COLUMNS if c[1] in (
        "buyer", "category", "status", "website", "tiers", "n_states", "states", "loan_min",
        "loan_max_label", "income_min", "filters_label", "notes")]
    wb_cols.insert(4, ("Most leads sent to one tier", "max_sent", "#,##0", 11, False))
    _sheet(wb, "Win-back", wb_cols, winback)

    issues = []
    for alias in unknown_aliases(conn):
        issues.append({"issue": "Name not in buyer reference file", "buyer": alias,
                       "detail": "Add it to data/reference/buyers.csv with a category."})
    for p in tiered:
        if not p["has_filters"]:
            issues.append({"issue": "No filters exported", "buyer": p["buyer"],
                           "detail": "Profile columns blank; will rely on website research."})
    for r in conn.execute(
            """SELECT COALESCE(b.canonical_name, f.alias) AS buyer, COUNT(*) AS n
               FROM tier_filter f LEFT JOIN buyer_alias a ON a.alias = f.alias
               LEFT JOIN buyer b ON b.id = a.buyer_id
               WHERE f.tier_name IS NULL GROUP BY 1 ORDER BY 1"""):
        issues.append({"issue": "Unnamed tier(s) in filter export", "buyer": r["buyer"],
                       "detail": f"{r['n']} filter rows had a value in the tier column; grouped as one unnamed tier."})
    _sheet(wb, "Data issues", [("Issue", "issue", None, 34, False), ("Buyer", "buyer", None, 30, False),
                               ("Detail", "detail", None, 80, False)], issues)

    facts = fact_rows(conn)
    if facts:
        _sheet(wb, "Research facts", [
            ("Buyer", "buyer", None, 28, False), ("Website", "website", None, 24, False),
            ("Field", "field", None, 18, False), ("Value", "value", None, 45, False),
            ("Source", "source_url", None, 45, False), ("Supporting text", "quote", None, 60, False),
            ("Verified by", "verified_by", None, 12, True)], facts)

    tier_rows = [dict(r) for r in conn.execute(
        """SELECT COALESCE(b.canonical_name, r.alias) AS buyer, r.tier_name, r.is_price_reject,
                  r.sent, r.total_sold, r.commission, r.epl, r.redirected
           FROM tier_revenue r LEFT JOIN buyer_alias a ON a.alias = r.alias
           LEFT JOIN buyer b ON b.id = a.buyer_id ORDER BY r.commission DESC""")]
    _sheet(wb, "Tiers", [
        ("Buyer", "buyer", None, 30, False), ("Tier", "tier_name", None, 36, False),
        ("Price reject tier", "is_price_reject", None, 8, False), ("Sent", "sent", "#,##0", 9, False),
        ("Accepted leads", "total_sold", "#,##0", 10, False), ("Revenue ($)", "commission", "#,##0", 11, False),
        ("EPL ($)", "epl", "#,##0.00", 9, False), ("Redirected", "redirected", "#,##0", 9, False)], tier_rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"buyers": len(profiles), "win_back": len(winback), "issues": len(issues)}


def write_targets(conn: sqlite3.Connection, path: Path) -> dict:
    from .scoring import WEIGHTS, score
    from .discover import universe_rows
    targets = score(conn)
    researched = [t for t in targets if t["researched"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Prospect targets", True),
        (f"Generated {date.today().isoformat()}. Universe: companies in the CFPB complaint database with installment, "
         "payday or personal line of credit complaints in the last 12 months.", False),
        ("", False),
        ("Sheets", True),
        ("Targets: researched prospects, scored. Shaded columns are for you (decision, owner, notes).", False),
        ("Universe: every company in the CFPB pull with its status (client, screened out, candidate).", False),
        ("", False),
        ("Score (max ~90)", True),
        (f"Size up to {WEIGHTS['size']} (complaint volume, log scale) + type fit up to {WEIGHTS['type']} "
         f"(operator/servicer > tribal > bank-partner > state-licensed) + online {WEIGHTS['online']} + lead-buying "
         f"evidence {WEIGHTS['lead_signals']} + segment (subprime +5, prime -20).", False),
        ("Fit: 'price' = installment / loans up to $2,500+ (like your high-revenue buyers); 'volume' = small-dollar "
         "(like your high-volume buyers).", False),
        ("", False),
        ("Caveats", True),
        ("Complaint counts are a rough size signal: larger lenders get more complaints, but so do badly run ones.", False),
        ("Research facts come from public pages and are cited on the Sources column; check before relying on them. "
         "Litigation is shown for awareness, not scored.", False),
        ("Existing clients were matched by hand (data/reference/company_matches.csv); a missed match can show a client "
         "as a target — tell me and I'll add it.", False),
        ("This file is confidential. Contacts are not included yet (step 3).", True),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)

    for t in researched:
        t["sources_text"] = "\n".join(sorted(t["sources"])[:6])
        t["decision"] = t.get("decision")
    cols = [
        ("Rank", "rank", None, 5, False),
        ("Company (CFPB name)", "company", None, 32, False),
        ("Website", "website", None, 24, False),
        ("Score", "score", "0.0", 7, False),
        ("Fit", "fit", None, 12, False),
        ("Type", "r_company_type", None, 22, False),
        ("Why it fits", "why", None, 70, False),
        ("Operator group", "operator_group", None, 28, False),
        ("Consumer brands", "r_consumer_brands", None, 40, False),
        ("Parent / servicer", "r_parent_or_servicer", None, 30, False),
        ("Tribe", "r_tribe", None, 22, False),
        ("Products", "r_products", None, 26, False),
        ("Loan range", "r_loan_amount_range", None, 20, False),
        ("States", "r_states_served", None, 30, False),
        ("Online / branch", "r_storefront_or_online", None, 14, False),
        ("Segment", "r_customer_segment", None, 12, False),
        ("Lead-buying evidence", "r_lead_buying_signals", None, 40, False),
        ("Litigation / regulatory", "r_litigation_or_regulatory", None, 40, False),
        ("CFPB complaints (12m)", "complaints_total", "#,##0", 10, False),
        ("Roles to approach", "roles", None, 40, False),
        ("Research confidence", "r_research_confidence", None, 10, False),
        ("Sources", "sources_text", None, 50, False),
        ("Score: size", "score_size", "0.0", 7, False),
        ("Score: type", "score_type", "0.0", 7, False),
        ("Score: online", "score_online", "0.0", 7, False),
        ("Score: lead signals", "score_lead_signals", "0.0", 7, False),
        ("Score: segment", "score_segment", "0", 7, False),
        ("Decision (pursue / park / drop)", "decision", None, 14, True),
        ("Owner", "owner", None, 12, True),
        ("Your notes", "your_notes", None, 40, True),
    ]
    for i, t in enumerate(researched, 1):
        t["rank"] = i
    _sheet(wb, "Targets", cols, researched, freeze="C2")

    uni = universe_rows(conn)
    for u in uni:
        u["status"] = ("existing client / group" if u["buyer"] else
                       f"screened out: {u['screened_out']}" if u.get("screened_out") else
                       "candidate - researched" if any(t["company"] == u["company"] and t["researched"]
                                                       for t in researched) else "candidate - not yet researched")
    _sheet(wb, "Universe", [
        ("Company (CFPB name)", "company", None, 40, False), ("Status", "status", None, 40, False),
        ("Matched buyer", "buyer", None, 26, False), ("Operator group", "operator_group", None, 30, False),
        ("Complaints (12m)", "complaints_total", "#,##0", 10, False),
        ("Installment", "complaints_installment_loan", None, 9, False),
        ("Payday", "complaints_payday_loan", None, 9, False),
        ("Line of credit", "complaints_personal_line_of_credit", None, 9, False)], uni)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"targets": len(researched), "universe": len(uni)}
