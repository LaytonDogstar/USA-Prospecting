"""End-to-end test on synthetic exports. Uses made-up buyer names only."""
import csv
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from prospecting import db
from prospecting.export import write_review
from prospecting.importers import import_filters, import_revenue, load_reference, unknown_aliases
from prospecting.names import looks_like_filter_value, parse_range, split_buyer_tier
from prospecting.profiles import build_profiles
from prospecting.research import fact_rows, load_research, research_by_buyer

HEADER = ["Name", "Processed", "Sent", "Unique Sold", "Multi Sell", "Total Sold", "Declined", "Error",
          "Redirected", "Remarketed", "Commission", "EPL", "Response Time", "Tree", "Tree %"]


def revenue_row(name, sold, commission, sent=1000, redirected=0):
    return [name, 2000, sent, sold, 0, sold, 0, 0, redirected, 0, commission,
            commission / sold if sold else 0, "0.10 secs", None, "0.00 %"]


class NamesTest(unittest.TestCase):
    def test_split_nested(self):
        self.assertEqual(split_buyer_tier("Acme Corp (Brand) (T1 Day (BR) - ID 104)"),
                         ("Acme Corp (Brand)", "T1 Day (BR) - ID 104"))
        self.assertEqual(split_buyer_tier("Plain Lender"), ("Plain Lender", ""))

    def test_ranges(self):
        self.assertEqual(parse_range("200.00 - 2000.00"), (200.0, 2000.0))
        self.assertEqual(parse_range("800.00 - 0.00"), (800.0, None))
        self.assertTrue(looks_like_filter_value("0.00 - 0.00"))
        self.assertFalse(looks_like_filter_value("T1 (Day)"))


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Admin"])
        ws.append(HEADER)
        rows = [
            revenue_row("Big Price (T1)", 100, 20000),        # high revenue, low volume
            revenue_row("Big Price (Price Reject Tier)", 10, 500),
            revenue_row("Big Volume (T1)", 5000, 5000),       # high volume, low price
            revenue_row("Big Volume V2 (T1)", 1000, 1000),    # alias merged into Big Volume
            revenue_row("Small Lender (T1)", 5, 50),
            revenue_row("Dormant Lender (T1)", 0, 0, sent=900),
            revenue_row("Some Network (10.00)", 9000, 90000),
        ]
        for r in rows + [rows[-1]]:                           # last row duplicated verbatim
            ws.append(r)
        self.revenue = self.tmp / "rev.xlsx"
        wb.save(self.revenue)

        self.filters = self.tmp / "filters.csv"
        with open(self.filters, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["Big Price", "T1", "State Blacklist", "NY, CA"])
            w.writerow(["Big Price", "T1", "Loan Amount", "500.00 - 5000.00"])
            w.writerow(["Big Price", "T2", "Loan Amount", "300.00 - 0.00"])   # T2: no blacklist, no max
            w.writerow(["Big Volume", "0.00 - 0.00", "Direct Deposit Only", "True"])
            w.writerow(["Big Volume", "800.00 - 0.00", "Monthly Income", "800.00 - 0.00"])
            w.writerow(["Big Volume", "0.00 - 0.00", "State Blacklist", "TX"])

        self.reference = self.tmp / "buyers.csv"
        with open(self.reference, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["alias", "canonical_name", "category", "status", "website", "notes"])
            w.writerow(["Big Price", "", "D", "", "", ""])
            w.writerow(["Big Volume", "", "D", "", "", ""])
            w.writerow(["Big Volume V2", "Big Volume", "", "", "", ""])
            w.writerow(["Small Lender", "", "D", "", "", ""])
            w.writerow(["Dormant Lender", "", "D", "", "", ""])
            w.writerow(["Some Network", "", "N", "", "", ""])

        self.conn = db.connect(self.tmp / "t.db")
        self.rev = import_revenue(self.conn, self.revenue, "test period")
        import_filters(self.conn, self.filters)
        load_reference(self.conn, self.reference)

    def profiles(self):
        return {p["buyer"]: p for p in build_profiles(self.conn)}

    def test_duplicates_skipped_and_aliases_merged(self):
        self.assertEqual(self.rev["duplicates_skipped"], 1)
        self.assertEqual(unknown_aliases(self.conn), [])
        self.assertEqual(self.profiles()["Big Volume"]["accepted"], 6000)

    def test_price_reject_revenue_counts(self):
        p = self.profiles()["Big Price"]
        self.assertEqual(p["revenue"], 20500)
        self.assertAlmostEqual(p["price_reject_share"], 500 / 20500)

    def test_filters_take_loosest_values(self):
        p = self.profiles()["Big Price"]
        self.assertEqual(len(p["states_accepted"]), 51)        # T2 has no blacklist
        self.assertEqual(p["loan_min"], 300)
        self.assertTrue(p["loan_max_unbounded"])
        v = self.profiles()["Big Volume"]
        self.assertNotIn("TX", v["states_accepted"])            # unnamed rows grouped as one tier
        self.assertEqual(v["income_min"], 800)

    def test_value_tiers(self):
        p = self.profiles()
        self.assertEqual(p["Big Price"]["price_tier"], "A")
        self.assertEqual(p["Big Price"]["reference_segment"], "price")
        self.assertEqual(p["Big Volume"]["reference_segment"], "volume")
        self.assertIsNone(p["Some Network"]["volume_tier"])     # networks excluded from tiering
        self.assertEqual(p["Dormant Lender"]["status"], "no sales this period")

    def test_research(self):
        path = self.tmp / "research.json"
        path.write_text(json.dumps([
            {"buyer": "Big Volume V2", "website": "bigvolume.example", "confidence": "high", "notes": "n",
             "facts": [{"field": "company_type", "value": "tribal", "source_url": "https://bigvolume.example/about",
                        "quote": "owned by the Example Tribe"},
                       {"field": "products", "value": ["installment loan", "line of credit"],
                        "source_url": "https://bigvolume.example/rates", "quote": "q"},
                       {"field": "hq_location", "value": "Somewhere", "source_url": None}]},
            {"buyer": "Nobody We Know", "website": None, "confidence": "low", "facts": []},
        ]))
        res = load_research(self.conn, [path])
        self.assertEqual(res, {"companies": 2, "facts": 2, "skipped_no_source": 1, "unmatched": 1})
        r = research_by_buyer(self.conn)["Big Volume"]          # alias resolved to canonical buyer
        self.assertEqual(r["r_company_type"], "tribal")
        self.assertEqual(r["r_products"], "installment loan, line of credit")
        self.assertEqual(len(fact_rows(self.conn)), 2)
        load_research(self.conn, [path])                        # reloading replaces, not duplicates
        self.assertEqual(len(fact_rows(self.conn)), 2)

    def test_workbook(self):
        out = self.tmp / "review.xlsx"
        res = write_review(self.conn, out, "test period")
        self.assertEqual(res["win_back"], 1)
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Read me", "Buyers", "Reference profile", "Win-back", "Data issues", "Tiers"])


if __name__ == "__main__":
    unittest.main()
